import asyncio
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pymupdf
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app.config import settings
from app.main import app
from app.pipeline.extract import extract_statements
from app.pipeline.runner import run_pipeline
from app.pipeline.statement_extractor import EvidenceReader, Locator, NAMES, Statement, TableExtractor, Writer

SAMPLE = Path(__file__).resolve().parent / "fixtures/extraction/inputs/Press release Carrefour Q4+FY 2025.pdf"


class ExtractionTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp = self.enterContext(tempfile.TemporaryDirectory())
        self.root = Path(self.temp)
        self.enterContext(patch.object(settings, "extraction_dir", self.root / "artifacts"))
        self.enterContext(patch.object(settings, "extract_use_llm", False))
        self.enterContext(patch.object(settings, "extract_ocr_mode", "off"))

    async def test_sophie_sample_through_real_extraction_and_downloads(self):
        emit = AsyncMock()
        payload = await extract_statements(SAMPLE, emit)
        expected = {"IncomeStatement": [18], "Balance Sheet": [19], "CashFlow": [20]}
        self.assertEqual(payload["missing_statements"], [])
        for name, pages in expected.items():
            self.assertEqual(payload["statements"][name]["source_pages"], pages)
            self.assertTrue(payload["statements"][name]["found"])
            self.assertGreater(len(payload["statements"][name]["rows"]), 2)
        self.assertEqual(emit.await_args_list[-1].args[1], "done")
        self.assertEqual(emit.await_args_list[-1].args[3], payload)
        with TestClient(app) as client:
            response = client.get(payload["artifacts"]["json"])
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json(), payload)
            workbook_response = client.get(payload["artifacts"]["xlsx"])
            self.assertEqual(workbook_response.status_code, 200)
        workbook = load_workbook(BytesIO(workbook_response.content))
        self.assertEqual(set(workbook.sheetnames), set(NAMES))
        self.assertIn("[20]", workbook["CashFlow"]["A2"].value)
        workbook.close()

    async def test_translated_text_and_missing_tables(self):
        report = self.root / "translated.txt"
        report.write_text("Consolidated income statement\n(in USD millions)\nItem  2025  2024\nSales  100  90\nProfit  20  10\nNote 1: unaudited", encoding="utf-8")
        payload = await extract_statements(report, AsyncMock())
        self.assertEqual(payload["missing_statements"], ["CashFlow", "Balance Sheet"])
        income = payload["statements"]["IncomeStatement"]
        self.assertTrue(income["found"])
        self.assertEqual(income["currency_and_scale"], "(in USD millions)")
        self.assertEqual(income["footnotes"], ["Note 1: unaudited"])
        self.assertIn(["Sales", "100", "90"], income["rows"])

    async def test_plain_prose_is_evidence_not_a_recovered_table(self):
        report = self.root / "narrative.txt"
        report.write_text("Our balance sheet remains strong.\nWe discuss performance here.", encoding="utf-8")
        payload = await extract_statements(report, AsyncMock())
        balance = payload["statements"]["Balance Sheet"]
        self.assertFalse(balance["found"])
        self.assertTrue(balance["text_evidence"])
        self.assertEqual(len(payload["missing_statements"]), 3)

    async def test_unreadable_input_emits_error_and_stops_pipeline(self):
        report = self.root / "empty.txt"
        report.write_text("", encoding="utf-8")
        emit = AsyncMock()
        with self.assertRaisesRegex(ValueError, "No readable report text"):
            await extract_statements(report, emit)
        self.assertEqual(emit.await_args_list[-1].args[1], "error")

    async def test_pipeline_passes_extraction_payload_to_final_event(self):
        payload = {"source_pdf": "report.pdf", "statements": {}, "artifacts": {"json": "/artifacts/test/report.json"}}
        with patch("app.pipeline.runner.find_statements", new_callable=AsyncMock,
                   return_value=[{"path": Path("report.pdf"), "year": 2025}]), \
             patch("app.pipeline.runner.check_language", new_callable=AsyncMock,
                   return_value={"language": "English", "is_english": True}), \
             patch("app.pipeline.runner.extract_statements", new_callable=AsyncMock, return_value=payload):
            emit = AsyncMock()
            await run_pipeline({"name": "Sample"}, "France", emit)
        final = emit.await_args_list[-1].args
        self.assertEqual(final[0:2], ("pipeline", "done"))
        self.assertEqual(final[3]["results"][0]["statements"], payload)


class LocatorTests(unittest.IsolatedAsyncioTestCase):
    async def test_llm_can_reject_a_heuristic_table(self):
        pages = ["Cash flow is discussed here. No table is present."]
        with patch.object(settings, "openrouter_api_key", "test"), \
             patch("app.pipeline.statement_extractor.chat_json", new_callable=AsyncMock,
                   return_value={name: [] for name in NAMES}) as ai:
            result = await Locator("test-model").locate(pages)
        self.assertEqual(result, {name: [] for name in NAMES})
        self.assertEqual(ai.call_args.args[0], "test-model")

    async def test_invalid_llm_pages_fall_back_with_warning(self):
        locator = Locator("test-model")
        with patch.object(settings, "openrouter_api_key", "test"), \
             patch("app.pipeline.statement_extractor.chat_json", new_callable=AsyncMock,
                   return_value={"CashFlow": [999], "Balance Sheet": [], "IncomeStatement": []}):
            result = await locator.locate(["Consolidated cash flow statement"])
        self.assertEqual(result["CashFlow"], [1])
        self.assertIn("heuristic", locator.warnings[0])

    async def test_cancellation_does_not_fall_back_to_heuristics(self):
        with patch.object(settings, "openrouter_api_key", "test"), \
             patch("app.pipeline.statement_extractor.chat_json", new_callable=AsyncMock,
                   side_effect=asyncio.CancelledError):
            with self.assertRaises(asyncio.CancelledError):
                await Locator("test-model").locate(["Consolidated cash flow statement"])


class ReaderWriterTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(self.enterContext(tempfile.TemporaryDirectory()))

    def scanned_pdf(self):
        pdf = self.root / "scan.pdf"
        image_data = BytesIO()
        with Image.new("RGB", (200, 100), "white") as image:
            image.save(image_data, format="PNG")
        with pymupdf.open() as doc:
            page = doc.new_page()
            page.insert_image(page.rect, stream=image_data.getvalue())
            doc.save(pdf)
        return pdf

    def test_ocr_text_is_used_by_table_extraction(self):
        pdf = self.scanned_pdf()
        with patch("app.pipeline.statement_extractor.shutil.which", return_value="tesseract"), \
             patch("app.pipeline.statement_extractor.pytesseract.image_to_string",
                   return_value="Consolidated income statement\nItem  2025  2024\nSales  100  90\nProfit  20  10"):
            evidence = EvidenceReader().read(pdf)
        self.assertEqual(evidence.ocr_pages, [1])
        with pymupdf.open(pdf) as doc:
            result = TableExtractor().extract(doc, evidence, "IncomeStatement", [1])
        self.assertTrue(result.found)
        self.assertIn(["Sales", "100", "90"], result.rows)
        self.assertTrue(any("OCR" in warning for warning in result.warnings))

    def test_required_ocr_without_tesseract_fails_explicitly(self):
        pdf = self.scanned_pdf()
        with patch("app.pipeline.statement_extractor.shutil.which", return_value=None):
            with self.assertRaisesRegex(RuntimeError, "Tesseract"):
                EvidenceReader("required").read(pdf)

    def test_workbook_stores_pdf_formula_text_without_executing_it(self):
        items = {name: Statement(name) for name in NAMES}
        items["CashFlow"] = Statement("CashFlow", found=True, rows=[["=1+2", "100"]])
        out = self.root / ("a" * 32)
        Writer().write(Path("report.v2.pdf"), out, items, [])
        workbook = load_workbook(out / "report.xlsx", data_only=False)
        self.assertEqual(workbook["CashFlow"]["A5"].value, "=1+2")
        self.assertEqual(workbook["CashFlow"]["A5"].data_type, "s")
        workbook.close()
        self.assertEqual(json.loads((out / "report.json").read_text())["source_pdf"], "report.v2.pdf")

    def test_artifact_endpoint_rejects_unknown_ids_and_files(self):
        with patch.object(settings, "extraction_dir", self.root), TestClient(app) as client:
            for url in ("/artifacts/not-an-id/report.json", f"/artifacts/{'a' * 32}/secret.env",
                        f"/artifacts/{'a' * 32}/report.json"):
                self.assertEqual(client.get(url).status_code, 404)


if __name__ == "__main__":
    unittest.main()
